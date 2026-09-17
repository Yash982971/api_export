import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_singing_bowl_pptx(output_path):
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625) # 16:9 widescreen
    
    blank_layout = prs.slide_layouts[6]
    
    # Colors
    gold = RGBColor(212, 175, 55)
    white = RGBColor(255, 255, 255)
    dark_bg = RGBColor(18, 14, 8)
    card_bg = RGBColor(30, 23, 16)
    
    # Slide 1: Cover
    slide1 = prs.slides.add_slide(blank_layout)
    txBox = slide1.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(3.625))
    tf = txBox.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "OM ENTERPRISE"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = gold
    p.alignment = PP_ALIGN.CENTER
    
    p2 = tf.add_paragraph()
    p2.text = "Authentic Handmade Himalayan Singing Bowls"
    p2.font.size = Pt(28)
    p2.font.bold = True
    p2.font.color.rgb = white
    p2.alignment = PP_ALIGN.CENTER
    
    p3 = tf.add_paragraph()
    p3.text = "Direct Manufacturer from Nepal | Wholesale • OEM • Private Label"
    p3.font.size = Pt(16)
    p3.font.color.rgb = gold
    p3.alignment = PP_ALIGN.CENTER
    
    # Slide 2: Product Range & Collections
    slide2 = prs.slides.add_slide(blank_layout)
    txBox2 = slide2.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(8.4), Inches(4.5))
    tf2 = txBox2.text_frame
    tf2.word_wrap = True
    
    hp = tf2.paragraphs[0]
    hp.text = "Product Range & Handcrafted Collections"
    hp.font.size = Pt(22)
    hp.font.bold = True
    hp.font.color.rgb = gold
    
    items = [
        "• Antique Finish Singing Bowls (Code: #HM-PLN) — Small 4\" to Extra-Large 30\" Yoga Bowls",
        "• Mantra & Buddha Hand Carved Bowls (Code: #HM-EC) — Tibetan Om Mani Padme Hum Engraving",
        "• Professional 7-Piece Chakra Sets (Code: #HM-7) — Tuned CDEFGAB with Mallets & Ring Cushions",
        "• Handmade Gongs — Plain & Mantra Carved (Code: #GNG) — 18\" and 24-26\" Sound Bath Gongs",
        "• Machine Made (Casted) Bowls & Accessories — Tingsha Cymbals, Meditation Bells & Happy Drums"
    ]
    for it in items:
        p = tf2.add_paragraph()
        p.text = it
        p.font.size = Pt(14)
        p.font.color.rgb = white
        
    # Slide 3: Contact & Export Partnership
    slide3 = prs.slides.add_slide(blank_layout)
    txBox3 = slide3.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(3.5))
    tf3 = txBox3.text_frame
    tf3.word_wrap = True
    
    cp = tf3.paragraphs[0]
    cp.text = "Start Your Export Partnership Today"
    cp.font.size = Pt(24)
    cp.font.bold = True
    cp.font.color.rgb = gold
    cp.alignment = PP_ALIGN.CENTER
    
    cp2 = tf3.add_paragraph()
    cp2.text = "OM ENTERPRISE — Direct Manufacturer from Nepal\n📧 exportindia2026us@gmail.com | 📱 +91 80577 10065 | 🌐 www.omenterprise.com"
    cp2.font.size = Pt(16)
    cp2.font.color.rgb = white
    cp2.alignment = PP_ALIGN.CENTER
    
    prs.save(output_path)
    print(f"PPTX presentation created at {output_path}")

if __name__ == '__main__':
    base_assets = os.path.join(os.path.dirname(__file__), 'assets')
    create_singing_bowl_pptx(os.path.join(base_assets, 'singing_bowl_product.pptx'))
